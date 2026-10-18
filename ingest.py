"""File parhna aur chhote chunks (page number ke sath) banana."""
import os
from pypdf import PdfReader
from docx import Document

CHUNK_WORDS = 150   # embedding model ~256 tokens tak parhta hai, is liye chhote chunks
OVERLAP = 30


def extract_pages(path):
    """[(page_number ya None, text), ...] wapas deta hai."""
    ext = path.lower().rsplit(".", 1)[-1]
    if ext in ("txt", "md"):
        with open(path, encoding="utf-8", errors="ignore") as f:
            return [(None, f.read())]
    if ext == "pdf":
        return [(i + 1, p.extract_text() or "") for i, p in enumerate(PdfReader(path).pages)]
    if ext == "docx":
        doc = Document(path)
        lines = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text for c in row.cells))
        return [(None, "\n".join(lines))]
    return []


def make_chunks(path):
    """[(chunk_text, page ya None), ...] - har chunk ke sath uska page number."""
    words = []
    for page, text in extract_pages(path):
        for w in text.split():
            words.append((w, page))
    step = CHUNK_WORDS - OVERLAP
    out = []
    for i in range(0, len(words), step):
        part = words[i:i + CHUNK_WORDS]
        if not part:
            break
        out.append((" ".join(w for w, _ in part), part[0][1]))
        if i + CHUNK_WORDS >= len(words):
            break
    return out


def main():
    """Command line se docs/ folder ki saari files embed karo: python ingest.py"""
    from rag import add_file
    base = os.path.dirname(os.path.abspath(__file__))
    docs = os.path.join(base, "docs")
    for name in sorted(os.listdir(docs)):
        n = add_file(os.path.join(docs, name), name)
        print(f"{name}: {n} chunks embed ho gaye")


if __name__ == "__main__":
    main()